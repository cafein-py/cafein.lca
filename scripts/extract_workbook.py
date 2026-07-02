"""Extract packaged datasets, mode defaults and golden test fixtures from the
pristine ITF "Good to go?" workbook.

Reads ``life-cycle-assessment-calculations-2020_original.xlsx`` (the unmodified
ITF download; the repo also holds a copy with a Finland 2020 electricity mix
edit which is NOT used here) and writes:

- ``cafein_lca/data/*.csv``      shared assumption tables (environment layer)
- ``cafein_lca/data/modes.csv``  per-mode-column scenario inputs + leaf data
- ``tests/data/golden.csv``      cached result values for all mode columns

Cell addresses in this script mirror the workbook; see plans/engine-notes.md
for the full formula map. Run from the repo root:

    python scripts/extract_workbook.py
"""

import csv
import pathlib
import re

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parent.parent
SOURCE = ROOT / "life-cycle-assessment-calculations-2020_original.xlsx"
DATA_DIR = ROOT / "cafein_lca" / "data"
TESTS_DATA_DIR = ROOT / "tests" / "data"

MODE_COL_FIRST, MODE_COL_LAST = 4, 137  # D..EG in 0_Total and stage sheets

MATERIALS = [
    "steel", "stainless_steel", "cast_iron", "wrought_aluminum",
    "cast_aluminum", "copper_brass", "glass", "plastics", "carbon_fiber",
    "rubber", "others",
]
DELIVERY_LEGS = [
    "air", "ship", "train_a", "train_b", "heavy_truck_a", "heavy_truck_b",
    "medium_truck_a", "medium_truck_b", "van_a", "van_b",
]
POWER_SOURCES = ["oil", "natural_gas", "coal", "nuclear", "biomass",
                 "other_renewables"]

# 0_Total result rows -> (metric, per, component) for the golden fixtures.
COMPONENTS = ["manufacturing", "delivery", "use", "services", "infrastructure"]
RESULT_BLOCKS = {
    "energy": {"pkm": (76, 77), "vkm": (83, 84), "vehicle": (90, 91)},
    "ghg": {"pkm": (98, 99), "vkm": (105, 106), "vehicle": (112, 113)},
}


def col_letter(idx):
    return openpyxl.utils.get_column_letter(idx)


def load(path):
    wf = openpyxl.load_workbook(path, read_only=True, data_only=False)
    wv = openpyxl.load_workbook(path, read_only=True, data_only=True)
    formulas, values = {}, {}
    for name in wf.sheetnames:
        formulas[name] = {
            (c.row, c.column): c.value
            for row in wf[name].iter_rows() for c in row if c.value is not None
        }
        values[name] = {
            (c.row, c.column): c.value
            for row in wv[name].iter_rows() for c in row if c.value is not None
        }
    return formulas, values


def write_csv(path, header, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)
    print(f"wrote {path.relative_to(ROOT)} ({len(rows)} rows)")


def extract_power_gen_mix(values):
    v = values["Power_Gen_Mix"]
    rows = []
    for col in range(2, 20):  # B..S
        region = v.get((3, col))
        if not region:
            continue
        rows.append([region] + [v.get((r, col), 0) for r in range(4, 10)])
    write_csv(DATA_DIR / "power_gen_mix.csv", ["region"] + POWER_SOURCES, rows)


def extract_fuel_properties(values):
    v = values["WTW_Fuel_properties"]
    # WTT energy factors (3_Use R27 VLOOKUP A4:C5).
    rows = [[v[(r, 1)], v[(r, 3)]] for r in (4, 5) if (r, 1) in v]
    write_csv(DATA_DIR / "fuel_wtt_energy.csv",
              ["fuel", "wtt_energy_factor"], rows)
    # Fuel GHG intensities (3_Use R31 / 4_Op R10 VLOOKUP A75:D90).
    rows = []
    for r in range(75, 91):
        if (r, 1) in v and v.get((r, 3)) is not None:
            rows.append([v[(r, 1)], v[(r, 3)], v.get((r, 4), "")])
    write_csv(DATA_DIR / "fuel_ghg.csv",
              ["fuel", "ghg_wtw_g_per_l", "ghg_ttw_g_per_l"], rows)
    # Electricity generation pathways (3_Use R28/R32: E45:E50, I45:I50).
    rows = []
    for i, r in enumerate(range(45, 51)):
        rows.append([POWER_SOURCES[i], v.get((r, 9)), v.get((r, 5))])
    write_csv(DATA_DIR / "electricity_pathways.csv",
              ["source", "wtt_energy_factor", "ghg_g_per_kwh"], rows)
    # Hydrogen pathways (3_Use R29/R33: names C57:C63, energy L, GHG X).
    rows = []
    for r in range(57, 64):
        name = v.get((r, 3))
        if name:
            rows.append([name, v.get((r, 12)), v.get((r, 24))])
    write_csv(DATA_DIR / "hydrogen_pathways.csv",
              ["pathway", "wtt_energy_factor", "ghg_g_per_mj"], rows)
    return {"electrolysis_label": v.get((64, 3))}


def extract_manufacturing_env(values):
    v = values["1_Manufacturing"]
    # Material production intensities by production region (EI=coal-Al, EJ=default).
    # Energy rows 38-51 map to: steel virgin/recycled, stainless, cast iron,
    # wAl virgin/recycled, cAl virgin/recycled, Cu, glass, plastics, CF,
    # rubber, others. GHG rows 86-99 mirror.
    names = [
        "virgin_steel", "recycled_steel", "stainless_steel", "cast_iron",
        "virgin_wrought_aluminum", "recycled_wrought_aluminum",
        "virgin_cast_aluminum", "recycled_cast_aluminum", "copper_brass",
        "glass", "plastics", "carbon_fiber", "rubber", "others",
    ]
    region_cols = {"With Al smelting mostly form coal": 139, "Default": 140}
    rows = []
    for i, name in enumerate(names):
        row = [name]
        for col in region_cols.values():
            row.append(v.get((38 + i, col)))
        for col in region_cols.values():
            row.append(v.get((86 + i, col)))
        rows.append(row)
    write_csv(
        DATA_DIR / "material_intensities.csv",
        ["material", "energy_mj_per_kg_coal_al", "energy_mj_per_kg_default",
         "ghg_g_per_kg_coal_al", "ghg_g_per_kg_default"], rows)

    # Assembly/disposal per kg (energy rows 55/56, GHG rows 103/104).
    rows = [
        ["assembly", v.get((55, 139)), v.get((55, 140)),
         v.get((103, 139)), v.get((103, 140))],
        ["disposal", v.get((56, 139)), v.get((56, 140)),
         v.get((104, 139)), v.get((104, 140))],
    ]
    write_csv(
        DATA_DIR / "assembly_disposal_intensities.csv",
        ["step", "energy_mj_per_kg_coal_al", "energy_mj_per_kg_default",
         "ghg_g_per_kg_coal_al", "ghg_g_per_kg_default"], rows)

    # Battery: specific energy (EI28:EQ29), manufacturing energy per vehicle
    # lifetime in mmBtu (EJ67:ER68) and GHG in g (EJ115:ER116), by chemistry
    # x region (row 67/115 = coal-Al, row 68/116 = default).
    rows = []
    for col in range(139, 148):  # EI..EQ chemistry headers row 28
        chem = v.get((28, col))
        if not chem:
            continue
        tcol = col + 1  # EJ..ER in the region x chemistry tables (headers row 66/114)
        rows.append([
            chem, v.get((29, col)),
            v.get((67, tcol)), v.get((68, tcol)),
            v.get((115, tcol)), v.get((116, tcol)),
        ])
    write_csv(
        DATA_DIR / "battery_chemistries.csv",
        ["chemistry", "specific_energy_kwh_per_kg",
         "mfg_energy_mmbtu_coal_al", "mfg_energy_mmbtu_default",
         "mfg_ghg_g_coal_al", "mfg_ghg_g_default"], rows)

    # Region labels used by 1_Manufacturing HLOOKUPs (row 37 headers).
    scalars = {
        "battery_assembly_energy_mj_per_kwh": v.get((60, 140)),
        "battery_disposal_energy_mj_per_kg": v.get((64, 140)),
        "battery_assembly_ghg_g_per_kwh": v.get((108, 140)),
        "battery_disposal_ghg_g_per_kg": v.get((112, 140)),
        "battery_reference_kwh": v.get((58, 72)),  # R58 col BT ([5]Battery_Sum!C22)
        "recycled_share_steel": v.get((22, 4)),
        "recycled_share_wrought_aluminum": v.get((23, 4)),
        "recycled_share_cast_aluminum": v.get((24, 4)),
        "mfg_region_coal_al": v.get((37, 139)),
        "mfg_region_default": v.get((37, 140)),
    }
    return scalars


def extract_delivery_env(values):
    v = values["2_Transport"]
    rows = []
    for i, leg in enumerate(DELIVERY_LEGS):
        rows.append([leg, v.get((59 + i, 4)), v.get((85 + i, 4))])
    write_csv(DATA_DIR / "delivery_legs.csv",
              ["leg", "energy_mj_per_tkm", "ghg_g_per_mj"], rows)


def extract_service_vehicles(values):
    v = values["4_Operational_Services"]
    rows = []
    for col in range(139, 152):  # EI..EU (EV handled below; range covers all)
        pass
    rows = []
    for col in range(139, 153):  # EI..EV
        name = v.get((3, col))
        if not name:
            continue
        rows.append([
            name,
            v.get((12, col), 0),
            v.get((15, col), 0),
        ])
    write_csv(DATA_DIR / "service_vehicles.csv",
              ["vehicle", "energy_mj_per_km", "ghg_g_per_km"], rows)


def extract_infrastructure_env(values):
    vi = values["Tech_Specs_Infra"]
    v5 = values["5_Infrastructure"]
    # Infrastructure types: materials t/km (A4:F9 cols C,E,F), annual use and
    # lifetime (A226:D231 cols B,C), share of energy imputable (A20:B25 col B).
    annual, imputable = {}, {}
    for r in range(226, 232):
        if (r, 1) in vi:
            annual[vi[(r, 1)]] = (vi.get((r, 2)), vi.get((r, 3)))
    for r in range(20, 26):
        if (r, 1) in vi:
            imputable[vi[(r, 1)]] = vi.get((r, 2))
    rows = []
    for r in range(4, 10):
        name = vi.get((r, 1))
        if not name:
            continue
        use, life = annual.get(name, (None, None))
        rows.append([name, vi.get((r, 3)), vi.get((r, 5)), vi.get((r, 6)),
                     life, use, imputable.get(name)])
    write_csv(
        DATA_DIR / "infrastructure_types.csv",
        ["infrastructure", "asphalt_t_per_km", "cement_t_per_km",
         "steel_t_per_km", "lifetime_years", "annual_use_mvkm",
         "energy_imputable_share"], rows)

    # Vehicle categories -> infrastructure assignment (A216:E222).
    rows = []
    for r in range(216, 223):
        name = vi.get((r, 1))
        if not name:
            continue
        rows.append([name, vi.get((r, 2)) or "", vi.get((r, 3)) or "",
                     vi.get((r, 4))])
    write_csv(DATA_DIR / "infrastructure_categories.csv",
              ["category", "infrastructure_1", "infrastructure_2",
               "share_infrastructure_1"], rows)

    scalars = {
        "infra_recycled_share_steel": v5.get((33, 4)),
        "infra_energy_asphalt_mj_per_kg": v5.get((36, 4)),
        "infra_energy_cement_mj_per_kg": v5.get((37, 4)),
        "infra_energy_virgin_steel_mj_per_kg": v5.get((38, 4)),
        "infra_energy_recycled_steel_mj_per_kg": v5.get((39, 4)),
        "infra_ghg_asphalt_g_per_kg": v5.get((59, 4)),
        "infra_ghg_cement_g_per_kg": v5.get((60, 4)),
        "infra_ghg_virgin_steel_g_per_kg": v5.get((61, 4)),
        "infra_ghg_recycled_steel_g_per_kg": v5.get((62, 4)),
        "reference_car_weight_kg": values["0_Total"].get((10, 41)),  # AO10
    }
    return scalars


def extract_constants(values, extra):
    vc = values["Convert"]
    consts = {
        "mj_per_kwh": vc.get((12, 2)),          # Convert!B12
        "mj_per_l_gasoline_x100": vc.get((17, 3)),  # Convert!C17 (per 100km use)
        "mj_per_mmbtu": vc.get((9, 7)),          # Convert!G9
    }
    consts.update(extra)
    rows = sorted(consts.items())
    write_csv(DATA_DIR / "constants.csv", ["name", "value"], rows)


FLUIDS_WEIGHT_SCALED = re.compile(r"C\$?AO\$?3")


def extract_modes(formulas, values):
    v0, f0 = values["0_Total"], formulas["0_Total"]
    vm, fm = values["1_Manufacturing"], formulas["1_Manufacturing"]
    vt, ft = values["2_Transport"], formulas["2_Transport"]
    vu = values["3_Use"]
    vo, fo = values["4_Operational_Services"], formulas["4_Operational_Services"]
    vi = values["5_Infrastructure"]

    header = (
        ["column", "name",
         # 0_Total scenario inputs
         "lifetime_years", "annual_km", "lifetime_km", "electricity_region",
         "vehicle_weight_kg", "battery_capacity_kwh", "production_region",
         "battery_chemistry", "hydrogen_pathway", "occupancy",
         "service_vehicle", "service_km_per_vehicle_day",
         "vehicles_per_service_trip",
         # manufacturing leaf
         ] + [f"share_{m}" for m in MATERIALS] + [
         "battery_replacements",
         "fluids_energy_mj", "fluids_ghg_g", "fluids_weight_scaled",
         # use leaf
         "fuel_consumption_per_100km", "electricity_consumption_kwh_per_km",
         "hydrogen_consumption_per_100km", "electric_driving_share",
         "fuel_type",
         # delivery leaf
         "delivery_weight_computed", "delivery_weight_kg",
         ] + [f"delivery_km_{leg}" for leg in DELIVERY_LEGS] + [
         # services leaf
         "services_uses_intensity_table",
         # infrastructure leaf (resolved per column)
         "infra_category", "infra1_type", "infra2_type",
         "infra1_asphalt_t_per_km", "infra1_cement_t_per_km",
         "infra1_steel_t_per_km", "infra2_asphalt_t_per_km",
         "infra2_cement_t_per_km", "infra2_steel_t_per_km",
         "infra1_lifetime_years", "infra2_lifetime_years",
         "infra1_annual_use_mvkm", "infra2_annual_use_mvkm",
         "infra_share_1", "infra_rail_allocation",
         ])

    rows = []
    for c in range(MODE_COL_FIRST, MODE_COL_LAST + 1):
        name = v0.get((2, c))
        if not name:
            continue
        letter = col_letter(c)

        def num(sheet_vals, r, col=c, default=0.0):
            val = sheet_vals.get((r, col))
            if val is None or isinstance(val, str):
                return default
            return val

        # 2_Transport R3: computed from weight+battery unless bespoke formula.
        t3_formula = ft.get((3, c))
        t3_canon = isinstance(t3_formula, str) and t3_formula.startswith(
            "=1_Manufacturing!"
        )
        # 4_Op R15: micromobility branch multiplies the intensity table.
        o15 = fo.get((15, c))
        uses_table = isinstance(o15, str) and "R12" not in o15 and (
            re.match(r"^=[A-Z]+\d*\$?12\b", "") or "12*" in o15.replace(" ", "")
        )
        # simpler and robust: branch A formulas look like '=C12*C10*C6'
        uses_table = isinstance(o15, str) and re.match(
            rf"^={letter}12\*{letter}10\*{letter}6$", o15.replace(" ", "")
        ) is not None

        f135 = fm.get((135, c))
        fluids_scaled = isinstance(f135, str) and bool(
            FLUIDS_WEIGHT_SCALED.search(f135.replace("$", ""))
            or "AO$3" in f135 or "$AO3" in f135 or "$AO$3" in f135
        )

        rail_alloc = c >= 128  # DX.. (rail/metro block: plain VLOOKUP share)

        row = (
            [letter, name,
             num(v0, 4), num(v0, 5), num(v0, 22), v0.get((8, c), ""),
             num(v0, 10), num(v0, 11), v0.get((12, c), ""),
             v0.get((13, c), ""), v0.get((14, c), ""), num(v0, 16),
             v0.get((19, c), ""), num(v0, 20), num(v0, 21),
             ] + [num(vm, 8 + i) for i in range(11)] + [
             num(vm, 34),
             num(vm, 135), num(vm, 137), int(fluids_scaled),
             num(vu, 3), num(vu, 4), num(vu, 5), num(vu, 6),
             vu.get((14, c), ""),
             int(t3_canon), num(vt, 3),
             ] + [num(vt, 7 + i) for i in range(10)] + [
             int(uses_table),
             vi.get((3, c), ""), vi.get((5, c), ""), vi.get((6, c), ""),
             num(vi, 11), num(vi, 12), num(vi, 13),
             num(vi, 16), num(vi, 17), num(vi, 18),
             num(vi, 21), num(vi, 22), num(vi, 25), num(vi, 26),
             num(vi, 7), int(rail_alloc),
             ])
        rows.append(row)
    write_csv(DATA_DIR / "modes.csv", header, rows)
    return [r[0] for r in rows]


def extract_golden(values):
    v0 = values["0_Total"]
    rows = []
    for c in range(MODE_COL_FIRST, MODE_COL_LAST + 1):
        name = v0.get((2, c))
        if not name:
            continue
        letter = col_letter(c)
        for metric, blocks in RESULT_BLOCKS.items():
            for per, (total_row, first_row) in blocks.items():
                vals = {"total": v0.get((total_row, c))}
                for i, comp in enumerate(COMPONENTS):
                    vals[comp] = v0.get((first_row + i, c))
                for comp, val in vals.items():
                    rows.append([letter, name, metric, per, comp,
                                 "" if val is None else val])
    write_csv(TESTS_DATA_DIR / "golden.csv",
              ["column", "name", "metric", "per", "component", "value"], rows)


def main():
    print(f"reading {SOURCE.name} ...")
    formulas, values = load(SOURCE)
    extract_power_gen_mix(values)
    wtw_extra = extract_fuel_properties(values)
    mfg_scalars = extract_manufacturing_env(values)
    extract_delivery_env(values)
    extract_service_vehicles(values)
    infra_scalars = extract_infrastructure_env(values)
    extra = {}
    extra.update(mfg_scalars)
    extra.update(infra_scalars)
    extra.update(wtw_extra)
    extract_constants(values, extra)
    cols = extract_modes(formulas, values)
    extract_golden(values)
    print(f"extracted {len(cols)} mode columns")


if __name__ == "__main__":
    main()
